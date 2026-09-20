"""Серверная генерация XLSX-спецификации (ТКП) через openpyxl.

Для крупных выгрузок вызывается из Celery (``app.tasks.jobs.export_quote_xlsx``);
для обычных размеров endpoint отдаёт файл синхронно."""

from __future__ import annotations

import uuid
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.enums import LicensingModel, QuoteLineKind
from app.models.quote import Quote
from app.services.xlsx_safe import safe_text

_KIND_LABELS = {
    QuoteLineKind.LICENSE: "Лицензия",
    QuoteLineKind.WORK: "Работы",
    QuoteLineKind.SUBCONTRACT: "Субподряд",
    QuoteLineKind.SUPPORT: "Техподдержка",
}
_MODEL_LABELS = {
    LicensingModel.PER_USER: "за пользователя",
    LicensingModel.PER_NAMED_USER: "именованный пользователь",
    LicensingModel.PER_CONCURRENT_USER: "конкурентный пользователь",
    LicensingModel.PER_CORE: "за ядро",
    LicensingModel.PER_CPU: "за CPU",
    LicensingModel.PER_SERVER: "за сервер",
    LicensingModel.PER_DEVICE: "за устройство",
    LicensingModel.SUBSCRIPTION: "подписка",
    LicensingModel.PERPETUAL: "бессрочно",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TOTAL_FONT = Font(bold=True)
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY_FMT = r"# ##0.00\ ₽"


def build_quote_xlsx_bytes(quote: Quote, project_code: str, project_name: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"ТКП v{quote.version}"

    ws["A1"] = safe_text(f"{quote.title} — {project_code}")
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = safe_text(project_name)
    ws["A3"] = f"Версия: {quote.version}    Буфер курса: {quote.currency_buffer_pct}%"
    ws["A3"].font = Font(italic=True, color="808080")

    headers = [
        "Тип",
        "Наименование",
        "Модель/метрика",
        "Кол-во",
        "Цена за ед.",
        "Валюта",
        "Скидка партн., %",
        "Скидка клиенту, %",
        "Себестоимость",
        "Цена продажи",
        "Маржа",
    ]
    header_row = 5
    for col, title in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=title)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER

    row = header_row + 1
    for line in sorted(quote.lines, key=lambda x: x.kind):
        model_metric = ""
        if line.licensing_model:
            model_metric = _MODEL_LABELS.get(line.licensing_model, str(line.licensing_model))
        if line.metric:
            model_metric = f"{model_metric} / {line.metric}" if model_metric else line.metric
        values = [
            _KIND_LABELS.get(line.kind, str(line.kind)),
            safe_text(line.name),
            safe_text(model_metric),
            float(line.qty),
            float(line.unit_price),
            line.currency,
            float(line.partner_discount_pct),
            float(line.client_discount_pct),
            float(line.cost_amount),
            float(line.sell_amount),
            float(line.margin),
        ]
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = _BORDER
            if col in (5, 9, 10, 11):
                c.number_format = _MONEY_FMT
        row += 1

    # Итоги.
    totals = quote.totals or {}
    row += 1
    ws.cell(row=row, column=8, value="ИТОГО:").font = _TOTAL_FONT
    for col, key in ((9, "total_cost"), (10, "total_sell"), (11, "margin")):
        c = ws.cell(row=row, column=col, value=float(Decimal(str(totals.get(key, "0")))))
        c.font = _TOTAL_FONT
        c.number_format = _MONEY_FMT
    ws.cell(
        row=row + 1,
        column=10,
        value=f"Маржинальность: {totals.get('margin_pct', '0')}%",
    ).font = _TOTAL_FONT

    widths = [14, 36, 22, 9, 14, 8, 14, 14, 16, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def build_quote_xlsx(session: AsyncSession, quote_id: uuid.UUID) -> str:
    """Версия для Celery: сохраняет файл во временный каталог, возвращает путь."""
    import tempfile

    from app.models.project import Project
    from app.repositories.quote_repo import QuoteRepository

    quote = await QuoteRepository(session).get_with_lines(quote_id)
    if quote is None:
        raise ValueError("Расчёт не найден")
    project = await session.get(Project, quote.project_id)
    data = build_quote_xlsx_bytes(
        quote, project.code if project else "", project.name if project else ""
    )
    path = f"{tempfile.gettempdir()}/quote_{quote.project_id}_v{quote.version}.xlsx"
    with open(path, "wb") as f:
        f.write(data)
    return path
