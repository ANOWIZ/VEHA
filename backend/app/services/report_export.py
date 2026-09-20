"""Серверная генерация XLSX-отчётов (utilization, таймшиты, портфель) через openpyxl.

Размеры выгрузок в прототипе небольшие — endpoint отдаёт файл синхронно (как и
экспорт ТКП). Для крупных периодов следует вынести в Celery."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.xlsx_safe import safe_text

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TOTAL_FONT = Font(bold=True)
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY_FMT = r"# ##0.00\ ₽"
_HOURS_FMT = r"# ##0.00"
_PCT_FMT = r"0.0"

_STAGE_LABELS = {
    "presale": "Пресейл",
    "survey": "Обследование",
    "design": "Проектирование",
    "implementation": "Внедрение",
    "pilot": "Опытная эксплуатация",
    "support": "Поддержка",
    "closed": "Закрыт",
}
_TS_STATUS_LABELS = {
    "draft": "Черновик",
    "submitted": "На утверждении",
    "approved": "Утверждён",
    "rejected": "Отклонён",
}
_RISK_LABELS = {"low_margin": "Низкая маржа", "hours_overrun": "Перерасход часов"}


def _write_header(ws: Worksheet, headers: list[str], row: int) -> None:
    for col, title in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=title)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER


def _autowidth(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_utilization_xlsx_bytes(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Загрузка ресурсов"

    ws["A1"] = "Отчёт по загрузке ресурсов (utilization)"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = (
        f"Период: {report['date_from']} — {report['date_to']}  ·  "
        f"рабочих дней: {report['working_days']}  ·  "
        f"ёмкость на сотрудника: {report['capacity_per_user']} ч"
    )
    ws["A2"].font = Font(italic=True, color="808080")

    headers = [
        "Сотрудник", "Подразделение", "Билируемые часы", "Ёмкость, ч", "Загрузка, %", "Проектов",
    ]
    hr = 4
    _write_header(ws, headers, hr)
    row = hr + 1
    for r in report["rows"]:
        values = [
            r["full_name"],
            r["department"] or "—",
            float(r["billable_hours"]),
            float(r["capacity_hours"]),
            float(r["utilization_pct"]),
            r["projects_count"],
        ]
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=safe_text(val))
            c.border = _BORDER
            if col in (3, 4):
                c.number_format = _HOURS_FMT
            elif col == 5:
                c.number_format = _PCT_FMT
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="ИТОГО:").font = _TOTAL_FONT
    tb = ws.cell(row=row, column=3, value=float(report["total_billable_hours"]))
    tb.font = _TOTAL_FONT
    tb.number_format = _HOURS_FMT
    tc = ws.cell(row=row, column=4, value=float(report["total_capacity_hours"]))
    tc.font = _TOTAL_FONT
    tc.number_format = _HOURS_FMT
    ta = ws.cell(row=row, column=5, value=float(report["avg_utilization_pct"]))
    ta.font = _TOTAL_FONT
    ta.number_format = _PCT_FMT

    _autowidth(ws, [32, 22, 18, 14, 14, 12])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_timesheet_xlsx_bytes(
    rows: list[dict], *, include_cost: bool, date_from: date, date_to: date
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Трудозатраты"

    ws["A1"] = "Выгрузка трудозатрат"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Период: {date_from} — {date_to}"
    ws["A2"].font = Font(italic=True, color="808080")

    headers = ["Дата", "Сотрудник", "Код", "Проект", "Задача", "Часы", "Статус", "Комментарий"]
    if include_cost:
        headers.append("Себестоимость")
    hr = 4
    _write_header(ws, headers, hr)

    row = hr + 1
    total_hours = Decimal("0")
    total_cost = Decimal("0")
    for r in rows:
        total_hours += r["hours"]
        values = [
            r["work_date"],
            r["user"],
            r["project_code"],
            r["project_name"],
            r["task"] or "—",
            float(r["hours"]),
            _TS_STATUS_LABELS.get(str(r["status"]), str(r["status"])),
            r["comment"],
        ]
        if include_cost:
            total_cost += r.get("cost", Decimal("0"))
            values.append(float(r.get("cost", 0)))
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=safe_text(val))
            c.border = _BORDER
            if col == 1:
                c.number_format = "yyyy-mm-dd"
            elif col == 6:
                c.number_format = _HOURS_FMT
            elif include_cost and col == 9:
                c.number_format = _MONEY_FMT
        row += 1

    row += 1
    ws.cell(row=row, column=5, value="ИТОГО:").font = _TOTAL_FONT
    th = ws.cell(row=row, column=6, value=float(total_hours))
    th.font = _TOTAL_FONT
    th.number_format = _HOURS_FMT
    if include_cost:
        tc = ws.cell(row=row, column=9, value=float(total_cost))
        tc.font = _TOTAL_FONT
        tc.number_format = _MONEY_FMT

    widths = [12, 28, 14, 34, 28, 10, 16, 40]
    if include_cost:
        widths.append(16)
    _autowidth(ws, widths)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_portfolio_xlsx_bytes(portfolio: dict, manager_names: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Портфель"

    ws["A1"] = "Портфель проектов: выручка, маржа, риски"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = (
        f"Проектов: {portfolio['projects_count']}  ·  "
        f"выручка: {portfolio['total_revenue']}  ·  "
        f"маржа: {portfolio['total_margin']} ({portfolio['avg_margin_pct']}%)  ·  "
        f"в зоне риска: {portfolio['at_risk']}"
    )
    ws["A2"].font = Font(italic=True, color="808080")

    headers = [
        "Код", "Проект", "Стадия", "РП", "Выручка", "Маржа", "Маржа, %",
        "Часы факт", "Часы план", "Перерасход, %", "Риски",
    ]
    hr = 4
    _write_header(ws, headers, hr)
    row = hr + 1
    for r in portfolio["rows"]:
        risks = ", ".join(_RISK_LABELS.get(x, x) for x in r["risks"]) or "—"
        values = [
            r["code"],
            r["name"],
            _STAGE_LABELS.get(str(r["stage"]), str(r["stage"])),
            manager_names.get(r["manager_id"], "—"),
            float(r["revenue"]),
            float(r["margin"]),
            float(r["margin_pct"]),
            float(r["actual_hours"]),
            float(r["planned_hours"]),
            float(r["hours_overrun_pct"]),
            risks,
        ]
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=safe_text(val))
            c.border = _BORDER
            if col in (5, 6):
                c.number_format = _MONEY_FMT
            elif col in (7, 10):
                c.number_format = _PCT_FMT
            elif col in (8, 9):
                c.number_format = _HOURS_FMT
        row += 1

    _autowidth(ws, [14, 34, 18, 24, 16, 16, 10, 12, 12, 14, 30])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
